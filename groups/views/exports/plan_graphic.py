# 📦 export_plan_graphic.py
# ️ Версия: b_0.0.4.14 (Versioned filenames + PermissionError handling)
# ✅ Статус: PRODUCTION-READY
# 📅 Последнее обновление: 2026-08-25

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from django.conf import settings
from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect
from groups.models import SchedulePlan
from reference.models import TrainingProgram, ProgramSubject
from datetime import datetime, timedelta
from pathlib import Path
from collections import defaultdict
import json
import re
import logging
from urllib.parse import quote

logger = logging.getLogger(__name__)

#  Принудительно русские месяца
RU_MONTHS = {
    1: 'Январь', 2: 'Февраль', 3: 'Март', 4: 'Апрель', 5: 'Май', 6: 'Июнь',
    7: 'Июль', 8: 'Август', 9: 'Сентябрь', 10: 'Октябрь', 11: 'Ноябрь', 12: 'Декабрь'
}


def _sanitize(name):
    if not name: return ""
    return re.sub(r'[<>:"/\\|?*\x00-\x1F]', '_', str(name).strip()).replace(' ', '_')


def export_plan_graphic_to_excel(request, plan_id):
    plan = get_object_or_404(SchedulePlan, pk=plan_id)
    group_category = getattr(plan.group, 'category', None)

    # --- 1. Сбор и подготовка данных ---
    class_days = plan.class_days or {}
    if isinstance(class_days, str):
        try:
            class_days = json.loads(class_days)
        except (json.JSONDecodeError, ValueError) as e:
            logger.warning("План-график pk=%s: некорректный JSON в class_days: %s", plan.pk, e)
            class_days = {}

    excluded_dates = plan.excluded_dates or []
    additional_dates = plan.additional_dates or []

    selected_program_id = request.GET.get('program_id')
    program = None
    if selected_program_id:
        try:
            program = TrainingProgram.objects.get(pk=selected_program_id)
        except TrainingProgram.DoesNotExist:
            logger.warning("План-график pk=%s: учебная программа id=%s не найдена", plan.pk, selected_program_id)

    if not program:
        if hasattr(plan, 'training_program') and plan.training_program:
            program = plan.training_program
        elif group_category:
            program = TrainingProgram.objects.filter(categories=group_category).first()

    program_subjects = []
    if program:
        exam_item, other_subjects = None, []
        for ps in ProgramSubject.objects.filter(program=program, is_enabled=True).select_related('subject').order_by(
                'subject__short_name'):
            item = {'code': ps.subject.short_name.lower(),
                    'display': ps.subject.short_name_display or ps.subject.short_name.upper()}
            if ps.subject.short_name.lower() == 'exam':
                exam_item = item
            else:
                other_subjects.append(item)
        program_subjects = other_subjects + ([exam_item] if exam_item else [])

    # 🔹 Коды предметов, которые реально выводятся в таблице
    subject_codes = {item['code'] for item in program_subjects}

    def _day_hours(day_data):
        """Сумма часов за день ТОЛЬКО по предметам программы (без служебных ключей)."""
        return sum(
            v for k, v in day_data.items()
            if (k in subject_codes if subject_codes else
                (not k.startswith('_') and not k.endswith('_topics') and k != 'scheduled'))
            and isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0
        )

    # Диапазон дат
    days_in_schedule = []
    current_date = plan.date_start
    while current_date <= plan.date_end:
        days_in_schedule.append(current_date)
        current_date += timedelta(days=1)

    # --- 2. ПРЕДВАРИТЕЛЬНЫЙ РАСЧЕТ СТРУКТУРЫ ТАБЛИЦЫ ---
    months_data = defaultdict(list)
    month_order = []
    days_with_classes = []

    for d in days_in_schedule:
        date_key = d.strftime('%Y-%m-%d')
        if date_key not in class_days: continue

        day_data = class_days[date_key]
        day_sum = _day_hours(day_data)

        if day_sum > 0:
            days_with_classes.append(d)
            m_label = f"{RU_MONTHS[d.month]} {d.year}"
            if m_label not in month_order: month_order.append(m_label)
            months_data[m_label].append(d)

    # 🔹 РАСЧЕТ ПОСЛЕДНЕЙ КОЛОНКИ ("ИТОГО")
    total_table_width_idx = 1 + len(days_with_classes) + 1
    total_col_letter = get_column_letter(total_table_width_idx)

    # Диапазон для объединения ячеек в шапке (УТВЕРЖДАЮ)
    merge_cols_count = 12
    merge_start_idx = max(2, total_table_width_idx - merge_cols_count)
    merge_start_letter = get_column_letter(merge_start_idx)
    merge_end_letter = total_col_letter

    # --- 3. Генерация Excel ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "План-график"

    font_small = Font(name='Arial Cyr', size=8)
    font_dates = Font(name='Arial Cyr', size=8, bold=True)
    font_bold = Font(name='Arial Cyr', size=10, bold=True)
    font_title = Font(name='Arial Cyr', size=14, bold=True)
    thin_border = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right_no_wrap = Alignment(horizontal='right', vertical='center', wrap_text=False)
    align_left_no_wrap = Alignment(horizontal='left', vertical='center', wrap_text=False)
    header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")

    # --- ШАПКА ---
    row = 1

    # Организация (слева)
    ws.cell(row=row, column=1, value=getattr(plan.group, 'organization_name', 'ООО "Своя автошкола"')).font = font_bold

    # 🔹 "УТВЕРЖДАЮ" - объединяем ячейки справа
    ws.merge_cells(f'{merge_start_letter}{row}:{merge_end_letter}{row}')
    cell = ws.cell(row=row, column=merge_start_idx, value='"УТВЕРЖДАЮ"')
    cell.font = font_bold
    cell.alignment = align_right_no_wrap
    ws.row_dimensions[row].height = None  # Авто-высота

    row += 1
    # "Директор" - объединяем ячейки справа
    ws.merge_cells(f'{merge_start_letter}{row}:{merge_end_letter}{row}')
    cell = ws.cell(row=row, column=merge_start_idx, value='Директор')
    cell.font = font_bold
    cell.alignment = align_right_no_wrap
    ws.row_dimensions[row].height = None  # Авто-высота

    row += 1
    # 🔹 Подпись - объединяем ячейки справа
    ws.merge_cells(f'{merge_start_letter}{row}:{merge_end_letter}{row}')
    cell = ws.cell(row=row, column=merge_start_idx, value='_______________ В.В. Евтушков')
    cell.font = font_bold
    cell.alignment = align_right_no_wrap
    ws.row_dimensions[row].height = None  # Авто-высота

    row += 1

    # 🔹 Дата - объединяем 8 ячеек справа, текст прижат к левому краю
    date_cols_count = 8
    date_merge_start_idx = max(2, total_table_width_idx - date_cols_count + 1)
    date_merge_start_letter = get_column_letter(date_merge_start_idx)

    ws.merge_cells(f'{date_merge_start_letter}{row}:{total_col_letter}{row}')
    ru_months_text = {1: 'января', 2: 'февраля', 3: 'марта', 4: 'апреля', 5: 'мая', 6: 'июня',
                      7: 'июля', 8: 'августа', 9: 'сентября', 10: 'октября', 11: 'ноября', 12: 'декабря'}
    date_text = f'"{plan.date_start.day:02d}" {ru_months_text[plan.date_start.month]} {plan.date_start.year} года'
    date_cell = ws.cell(row=row, column=date_merge_start_idx, value=date_text)
    date_cell.font = font_bold
    date_cell.alignment = align_left_no_wrap
    ws.row_dimensions[row].height = None  # Авто-высота

    row += 2

    # Заголовок План-графика (Строка 6)
    base_title = "План-график"
    if program and program.plan_graphic_title: base_title = program.plan_graphic_title

    ws.merge_cells(f'A{row}:{total_col_letter}{row}')
    title_cell = ws.cell(row=row, column=1, value=base_title)
    title_cell.font = font_title
    title_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = None  # 🔥 Автоматическая высота под 1–3 строки

    row += 2

    # Группа - на всю ширину таблицы
    group_num = getattr(plan.group, 'group_number', plan.group.id)
    ws.merge_cells(f'A{row}:{total_col_letter}{row}')
    ws.cell(row=row, column=1, value=f'учебная группа № {group_num}').font = font_bold
    ws.cell(row=row, column=1).alignment = align_center
    ws.row_dimensions[row].height = None
    row += 2

    # Информация (Преподаватель и т.д.) - ограничена шириной таблицы
    def format_dates_list(dates_list):
        if not dates_list:
            return "—"
        formatted = []
        for d in dates_list:
            try:
                val = datetime.strptime(d, '%Y-%m-%d').strftime('%d.%m.%Y') if isinstance(d, str) else d.strftime('%d.%m.%Y')
                formatted.append(val)
            except (ValueError, TypeError) as e:
                logger.warning("План-график pk=%s: пропущена некорректная дата %r: %s", plan.pk, d, e)
        return ', '.join(formatted) if formatted else "—"

    info_data = [
        ('Преподаватель', str(plan.teacher) if plan.teacher else '—'),
        ('Дни', str(plan.get_schedule_type_display())),
        ('Кроме', format_dates_list(excluded_dates)),
        ('Дополнительные дни', format_dates_list(additional_dates)),
        ('Время', f'{getattr(plan, "time_start", "09:00")} - {getattr(plan, "time_end", "17:00")}'),
        ('Место', getattr(plan, 'location', '—') or '—'),
    ]

    for label, value in info_data:
        ws.cell(row=row, column=1, value=label).font = font_bold
        ws.merge_cells(f'B{row}:{total_col_letter}{row}')
        val_cell = ws.cell(row=row, column=2, value=value)
        val_cell.font = font_small
        val_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.row_dimensions[row].height = None
        row += 1
    row += 1

    # --- ТАБЛИЦА ---
    start_col = 1
    ws.cell(row=row, column=start_col, value='Дата/Тема').font = font_bold
    ws.cell(row=row, column=start_col).border = thin_border
    ws.cell(row=row, column=start_col).alignment = align_center

    current_col = start_col + 1
    col_map = {}

    for m_label in month_order:
        days = months_data[m_label]
        ws.merge_cells(start_row=row, start_column=current_col, end_row=row, end_column=current_col + len(days) - 1)
        month_cell = ws.cell(row=row, column=current_col, value=m_label)
        month_cell.font = font_bold
        month_cell.alignment = align_center
        month_cell.border = thin_border
        current_col += len(days)

    total_col_idx = current_col
    ws.cell(row=row, column=total_col_idx, value='ИТОГО').font = font_bold
    ws.cell(row=row, column=total_col_idx).alignment = align_center
    ws.cell(row=row, column=total_col_idx).border = thin_border
    row += 1

    # 🔹 Дни (выделены жирным шрифтом)
    current_col = start_col + 1
    for m_label in month_order:
        for d in months_data[m_label]:
            col_map[d] = current_col
            cell = ws.cell(row=row, column=current_col, value=d.day)
            cell.font = font_dates
            cell.alignment = align_center
            cell.border = thin_border
            current_col += 1
    ws.cell(row=row, column=total_col_idx, value='').border = thin_border
    row += 1

    # Всего часов (считается ТОЛЬКО по предметам программы — сходится с суммой строк)
    ws.cell(row=row, column=start_col, value='Всего').font = font_bold
    ws.cell(row=row, column=start_col).border = thin_border
    ws.cell(row=row, column=start_col).fill = header_fill

    for d in days_with_classes:
        if d in col_map:
            c_idx = col_map[d]
            day_sum = _day_hours(class_days.get(d.strftime('%Y-%m-%d'), {}))
            cell = ws.cell(row=row, column=c_idx, value=int(day_sum) if day_sum == int(day_sum) else day_sum)
            cell.font = font_small
            cell.alignment = align_center
            cell.border = thin_border
            cell.fill = header_fill
    ws.cell(row=row, column=total_col_idx, value='').border = thin_border
    row += 1

    # Предметы
    for item in program_subjects:
        code, display_name = item['code'], item['display']
        ws.cell(row=row, column=start_col, value=display_name).font = font_small
        ws.cell(row=row, column=start_col).border = thin_border

        subject_total = 0
        for d in days_with_classes:
            if d in col_map:
                c_idx = col_map[d]
                h = class_days.get(d.strftime('%Y-%m-%d'), {}).get(code, 0)
                if not isinstance(h, (int, float)) or isinstance(h, bool):
                    h = 0

                cell = ws.cell(row=row, column=c_idx, value=int(h) if h > 0 and h == int(h) else (h if h > 0 else ''))
                cell.font = font_small
                cell.alignment = align_center
                cell.border = thin_border
                subject_total += h

        cell = ws.cell(row=row, column=total_col_idx,
                       value=int(subject_total) if subject_total == int(subject_total) else subject_total)
        cell.font = font_small
        cell.alignment = align_center
        cell.border = thin_border
        cell.fill = header_fill
        row += 1

    # --- Настройки печати ---
    ws.column_dimensions[get_column_letter(start_col)].width = 18
    for d in days_with_classes:
        if d in col_map: ws.column_dimensions[get_column_letter(col_map[d])].width = 3.5
    ws.column_dimensions[get_column_letter(total_col_idx)].width = 8

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = 0.3

    # --- Имя файла ---
    t_name = _sanitize(str(plan.teacher)) if plan.teacher else "Без_преподавателя"
    location_raw = getattr(plan, 'location', '') or 'без_адреса'
    location_clean = _sanitize(location_raw) or 'без_адреса'
    sched_val = getattr(plan, 'schedule_type', None) or getattr(plan.group, 'schedule_type', None)
    sched_map = {'even': 'четная', 'odd': 'нечетная', 'evening': 'вечерняя', '1': 'четная', '2': 'нечетная',
                 '3': 'вечерняя'}
    sched_name = sched_map.get(str(sched_val).lower().strip(), "четная")

    filename = f"План-график_гр{_sanitize(group_num)}_{t_name}_{sched_name}_{location_clean}.xlsx"

    # 🔹 Путь к папке экспорта берём из settings.py
    group_folder = Path(settings.EXPORT_DIR) / str(group_num)
    group_folder.mkdir(parents=True, exist_ok=True)
    base_path = group_folder / filename

    # 🔹 Сохранение: файл свободен — пишем как обычно;
    #    файл занят (открыт в Excel) — создаём копию с суффиксом (1), (2), ...
    saved_path = None
    try:
        wb.save(base_path)
        saved_path = base_path
    except PermissionError:
        logger.warning("Файл %s занят (открыт в Excel) — сохраняем копию с суффиксом", base_path)
        for i in range(1, 50):
            candidate = group_folder / f"{base_path.stem} ({i}){base_path.suffix}"
            try:
                wb.save(candidate)
                saved_path = candidate
                break
            except PermissionError:
                continue

    if saved_path is None:
        logger.error("План-график pk=%s: не удалось сохранить — все копии файла заняты", plan.pk)
        messages.error(request, '❌ Не удалось сохранить файл: все копии заняты. Закройте Excel и попробуйте снова.')
        return redirect(request.META.get('HTTP_REFERER', f'/groups/schedules/{plan.pk}/step2/'))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename*=UTF-8\'\'{quote(saved_path.name)}'
    with open(saved_path, "rb") as f:
        response.write(f.read())

    if saved_path != base_path:
        messages.warning(request, f'⚠️ Исходный файл занят, сохранена копия: {saved_path.name}')
    else:
        messages.success(request, f'✅ План-график экспортирован: {saved_path.name}')
    return response