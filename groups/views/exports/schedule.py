import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from groups.models import SchedulePlan
from reference.models import (
    SubjectDictionary,
    LessonTopic,
    ProgramTopic,
    TrainingProgram,
    ProgramSubject
)
from datetime import datetime, timedelta
import json
import re
from urllib.parse import quote


def _sanitize(name):
    """Очищает строку от недопустимых символов для имени файла"""
    if not name:
        return ""
    return re.sub(r'[<>:"/\\|?*\x00-\x1F]', '_', str(name).strip()).replace(' ', '_')


def export_schedule_to_excel(request, plan_id):
    plan = get_object_or_404(SchedulePlan, pk=plan_id)
    group_category = getattr(plan.group, 'category', None)

    subjects_map = {}
    if group_category:
        program = TrainingProgram.objects.filter(categories=group_category).first()
        if program:
            for ps in ProgramSubject.objects.filter(program=program, is_enabled=True).select_related('subject'):
                code = ps.subject.short_name.lower()
                display = ps.subject.short_name_display or ps.subject.short_name.upper()
                subjects_map[code] = display

    topics_content_map = {}
    if group_category:
        program = TrainingProgram.objects.filter(categories=group_category).first()
        if program:
            for ps in ProgramSubject.objects.filter(program=program, is_enabled=True).select_related('subject'):
                subject_code = ps.subject.short_name.lower()
                for pt in ProgramTopic.objects.filter(program_subject=ps).select_related('topic'):
                    topic = pt.topic
                    if subject_code not in topics_content_map:
                        topics_content_map[subject_code] = {}
                    topics_content_map[subject_code][topic.id] = {
                        'num': topic.topic_number,
                        'content': topic.content
                    }

    raw_time = getattr(plan, 'time_start', None) or '09:00'
    hour = int(raw_time.split(':')[0]) if isinstance(raw_time, str) else getattr(raw_time, 'hour', 9)
    time_type = "утро" if hour < 12 else ("день" if hour < 17 else "вечер")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Расписание"

    bold_font = Font(bold=True, size=12, name='Arial Cyr')
    title_font = Font(bold=True, size=14, name='Arial Cyr')
    thin_border = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='top', wrap_text=True)

    row = 1
    ws.merge_cells(f'A{row}:D{row}')
    ws[f'A{row}'] = '"СОСТАВИЛ и УТВЕРЖДАЮ"'
    ws[f'A{row}'].font = bold_font
    ws.row_dimensions[row].height = 20
    row += 1

    org_name = getattr(plan.group, 'organization_name', 'ООО "Своя автошкола"')
    ws[f'A{row}'] = f'Директор {org_name}'
    ws[f'A{row}'].font = bold_font
    ws.row_dimensions[row].height = 20
    row += 1

    ws[f'A{row}'] = '_______________В.В. Евтушков'
    ws[f'A{row}'].font = bold_font
    ws.row_dimensions[row].height = 20
    row += 1

    ru_months = {
        1: 'января', 2: 'февраля', 3: 'марта', 4: 'апреля', 5: 'мая', 6: 'июня',
        7: 'июля', 8: 'августа', 9: 'сентября', 10: 'октября', 11: 'ноября', 12: 'декабря'
    }
    ws[f'A{row}'] = f'"{plan.date_start.day:02d}" {ru_months[plan.date_start.month]} {plan.date_start.year} года'
    ws[f'A{row}'].font = bold_font
    row += 2

    group_num = getattr(plan.group, 'group_number', plan.group.id)
    titles = [
        ('РАСПИСАНИЕ', title_font),
        (f'занятий учебной группы № {group_num}', bold_font),
        ('в рамках оказания услуги по подготовке', bold_font),
        (f'водителей механических транспортных средств категории "{group_category.code if group_category else "B"}"',
         bold_font)
    ]
    for title_text, font in titles:
        ws.merge_cells(f'A{row}:F{row}')
        ws[f'A{row}'] = title_text
        ws[f'A{row}'].font = font
        ws[f'A{row}'].alignment = center_align
        ws.row_dimensions[row].height = 20
        row += 1
    row += 1

    headers = ['Дата', '', 'Кол-во часов', 'Предмет и краткое содержание занятий', 'Кто проводит занятия',
               'Место занятий']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[row].height = 40
    row += 1

    class_days = plan.class_days or {}
    if isinstance(class_days, str):
        try:
            class_days = json.loads(class_days)
        except:
            class_days = {}

    topics_data = {}
    for date_str, day_data in class_days.items():
        if date_str.startswith('_') or not isinstance(day_data, dict): continue
        for key, value in day_data.items():
            if key.endswith('_topics') and isinstance(value, dict):
                topics_data.setdefault(date_str, {})[key.replace('_topics', '')] = value

    current_date = plan.date_start
    while current_date <= plan.date_end:
        date_str = current_date.strftime('%Y-%m-%d')
        if date_str in class_days and not date_str.startswith('_'):
            day_data = class_days[date_str]
            for subject_code, hours in day_data.items():
                if subject_code.startswith('_') or subject_code.endswith('_topics'): continue
                if not isinstance(hours, (int, float)) or hours == 0: continue
                if subject_code == 'scheduled': continue

                display_name = subjects_map.get(subject_code.lower(), subject_code.upper())

                c1 = ws.cell(row=row, column=1, value=current_date.strftime('%d.%m.%Y'))
                c1.border = thin_border;
                c1.alignment = center_align
                c2 = ws.cell(row=row, column=2, value=display_name)
                c2.border = thin_border;
                c2.alignment = center_align
                c3 = ws.cell(row=row, column=3, value=int(hours))
                c3.border = thin_border;
                c3.alignment = center_align

                topics_text = ""
                if date_str in topics_data and subject_code in topics_data[date_str]:
                    topic_items = topics_data[date_str][subject_code]
                    topic_strings = []
                    subject_topics = topics_content_map.get(subject_code.lower(), {})
                    for tid, _ in topic_items.items():
                        try:
                            topic_pk = int(tid)
                            if topic_pk in subject_topics:
                                td = subject_topics[topic_pk]
                                topic_strings.append(f"Тема №{td['num']} - {td['content']}")
                        except:
                            pass
                    topics_text = "; ".join(topic_strings)

                c4 = ws.cell(row=row, column=4, value=topics_text)
                c4.border = thin_border;
                c4.alignment = left_align

                if subject_code.lower() == 'med' and plan.med_teacher:
                    teacher_name = str(plan.med_teacher)
                elif plan.teacher:
                    teacher_name = str(plan.teacher)
                else:
                    teacher_name = ''
                c5 = ws.cell(row=row, column=5, value=teacher_name)
                c5.border = thin_border;
                c5.alignment = center_align

                location = getattr(plan, 'location', '') or ''
                c6 = ws.cell(row=row, column=6, value=location)
                c6.border = thin_border;
                c6.alignment = center_align

                ws.row_dimensions[row].height = 35
                row += 1
        current_date += timedelta(days=1)

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 13
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 65
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 20

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = 0.4

    t_name = _sanitize(str(plan.teacher)) if plan.teacher else "Без_преподавателя"
    location_raw = getattr(plan, 'location', '') or 'без_адреса'
    location_clean = _sanitize(location_raw) or 'без_адреса'

    sched_val = getattr(plan, 'schedule_type', None) or getattr(plan.group, 'schedule_type', None)
    sched_map = {
        'even': 'четная', 'odd': 'нечетная', 'evening': 'вечерняя',
        '1': 'четная', '2': 'нечетная', '3': 'вечерняя'
    }
    sched_name = sched_map.get(str(sched_val).lower().strip(), "четная")

    filename_base = f"Расписание_гр{_sanitize(group_num)}_{t_name}_{time_type}_{sched_name}_{location_clean}"
    filename = f"{filename_base}.xlsx"

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    encoded_filename = quote(filename)
    ascii_filename = filename.encode('ascii', 'ignore').decode('ascii')
    response['Content-Disposition'] = f'attachment; filename*=UTF-8\'\'{encoded_filename}; filename="{ascii_filename}"'

    wb.save(response)
    return response