import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from groups.models import SchedulePlan
from reference.models import (
    SubjectDictionary,
    LessonTopic,
    ProgramTopic,
    TrainingProgram,
    ProgramSubject
)
from datetime import datetime, timedelta
import json
import re
from urllib.parse import quote


def _sanitize(name):
    """Очищает строку от недопустимых символов для имени файла"""
    if not name:
        return ""
    return re.sub(r'[<>:"/\\|?*\x00-\x1F]', '_', str(name).strip()).replace(' ', '_')


def export_schedule_to_excel(request, plan_id):
    plan = get_object_or_404(SchedulePlan, pk=plan_id)

    # 🔹 ЗАГРУЗКА ПРЕДМЕТОВ ЧЕРЕЗ ПРОГРАММУ ОБУЧЕНИЯ
    subjects_map = {}
    group_category = getattr(plan.group, 'category', None)

    if group_category:
        # 🔹 Находим программу обучения для этой категории
        program = TrainingProgram.objects.filter(
            categories=group_category
        ).first()

        if program:
            #  Загружаем предметы из программы
            program_subjects = ProgramSubject.objects.filter(
                program=program,
                is_enabled=True
            ).select_related('subject')

            for ps in program_subjects:
                subject = ps.subject
                code = subject.short_name.lower()
                display = subject.short_name_display if subject.short_name_display else subject.short_name.upper()
                subjects_map[code] = display

    # 🔹 Загружаем темы через программу обучения
    topics_content_map = {}

    if group_category:
        program = TrainingProgram.objects.filter(categories=group_category).first()

        if program:
            # 🔹 Загружаем все ProgramSubject для программы
            program_subjects = ProgramSubject.objects.filter(
                program=program,
                is_enabled=True
            ).select_related('subject')

            for ps in program_subjects:
                subject_code = ps.subject.short_name.lower()

                # 🔹 Загружаем темы для этого предмета в рамках программы
                program_topics = ProgramTopic.objects.filter(
                    program_subject=ps
                ).select_related('topic')

                for pt in program_topics:
                    topic = pt.topic
                    topic_pk = topic.id
                    topic_num = topic.topic_number
                    topic_content = topic.content

                    if subject_code not in topics_content_map:
                        topics_content_map[subject_code] = {}
                    topics_content_map[subject_code][topic_pk] = {
                        'num': topic_num,
                        'content': topic_content
                    }

    # 🔹 БЕЗОПАСНОЕ ПОЛУЧЕНИЕ ВРЕМЕНИ
    raw_time = getattr(plan, 'time_start', None)
    if not raw_time:
        raw_time = '09:00'

    if isinstance(raw_time, str):
        hour = int(raw_time.split(':')[0])
    elif hasattr(raw_time, 'hour'):
        hour = raw_time.hour
    else:
        hour = 9

    time_type = "утро" if hour < 12 else ("день" if hour < 17 else "вечер")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Расписание"

    bold_font = Font(bold=True, size=12, name='Arial Cyr')
    title_font = Font(bold=True, size=14, name='Arial Cyr')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # === ШАПКА ===
    row = 1
    ws.merge_cells(f'A{row}:D{row}')
    ws[f'A{row}'] = '"СОСТАВИЛ и УТВЕРЖДАЮ"'
    ws[f'A{row}'].font = bold_font
    ws.row_dimensions[row].height = 20
    row += 1

    org_name = getattr(plan.group, 'organization_name', 'ООО "Своя автошкола"')
    ws[f'A{row}'] = f'Директор {org_name}'
    ws[f'A{row}'].font = bold_font
    ws.row_dimensions[row].height = 20
    row += 1

    ws[f'A{row}'] = '_______________В.В. Евтушков'
    ws[f'A{row}'].font = bold_font
    ws.row_dimensions[row].height = 20
    row += 1

    ru_months = {
        1: 'января', 2: 'февраля', 3: 'марта', 4: 'апреля',
        5: 'мая', 6: 'июня', 7: 'июля', 8: 'августа',
        9: 'сентября', 10: 'октября', 11: 'ноября', 12: 'декабря'
    }
    day = plan.date_start.day
    month = ru_months[plan.date_start.month]
    year = plan.date_start.year
    ws[f'A{row}'] = f'"{day:02d}" {month} {year} года'
    ws[f'A{row}'].font = bold_font
    row += 2

    # === ЗАГОЛОВОК ===
    ws.merge_cells(f'A{row}:F{row}')
    ws[f'A{row}'] = 'РАСПИСАНИЕ'
    ws[f'A{row}'].font = title_font
    ws[f'A{row}'].alignment = center_alignment
    ws.row_dimensions[row].height = 20
    row += 1

    ws.merge_cells(f'A{row}:F{row}')
    group_num = getattr(plan.group, 'group_number', plan.group.id)
    ws[f'A{row}'] = f'занятий учебной группы № {group_num}'
    ws[f'A{row}'].font = bold_font
    ws[f'A{row}'].alignment = center_alignment
    row += 1

    ws.merge_cells(f'A{row}:F{row}')
    ws[f'A{row}'] = 'в рамках оказания услуги по подготовке'
    ws[f'A{row}'].font = bold_font
    ws[f'A{row}'].alignment = center_alignment
    row += 1

    ws.merge_cells(f'A{row}:F{row}')
    category = group_category.code if group_category else 'B'
    ws[f'A{row}'] = f'водителей механических транспортных средств категории "{category}"'
    ws[f'A{row}'].font = bold_font
    ws[f'A{row}'].alignment = center_alignment
    row += 2

    # === ТАБЛИЦА ===
    headers = ['Дата', '', 'Кол-во часов', 'Предмет и краткое содержание занятий', 'Кто проводит занятия',
               'Место занятий']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = bold_font
        cell.alignment = center_alignment
        cell.border = thin_border
    ws.row_dimensions[row].height = 40
    row += 1

    class_days = plan.class_days or {}
    if isinstance(class_days, str):
        try:
            class_days = json.loads(class_days)
        except Exception:
            class_days = {}

    topics_data = {}
    for date_str, day_data in class_days.items():
        if date_str.startswith('_') or not isinstance(day_data, dict):
            continue
        for key, value in day_data.items():
            if key.endswith('_topics') and isinstance(value, dict):
                subject = key.replace('_topics', '')
                topics_data.setdefault(date_str, {})[subject] = value

    # === ГЕНЕРАЦИЯ СТРОК ===
    current_date = plan.date_start
    while current_date <= plan.date_end:
        date_str = current_date.strftime('%Y-%m-%d')

        if date_str in class_days and not date_str.startswith('_'):
            day_data = class_days[date_str]

            for subject_code, hours in day_data.items():
                if subject_code.startswith('_') or subject_code.endswith('_topics'):
                    continue
                if not isinstance(hours, (int, float)) or hours == 0:
                    continue
                if subject_code == 'scheduled':
                    continue

                # 🔹 Получаем отображаемое название
                display_name = subjects_map.get(subject_code.lower(), subject_code.upper())

                ws.cell(row=row, column=1, value=current_date.strftime('%d.%m.%Y')).border = thin_border
                ws.cell(row=row, column=2, value=display_name).border = thin_border

                hours_cell = ws.cell(row=row, column=3, value=int(hours))
                hours_cell.border = thin_border
                hours_cell.alignment = center_alignment

                # 🔹 Формируем текст с темами
                topics_text = ""
                if date_str in topics_data and subject_code in topics_data[date_str]:
                    topic_items = topics_data[date_str][subject_code]
                    topic_strings = []
                    subject_topics = topics_content_map.get(subject_code.lower(), {})

                    for tid, topic_hours in topic_items.items():
                        try:
                            topic_pk = int(tid)
                            if topic_pk in subject_topics:
                                topic_data = subject_topics[topic_pk]
                                topic_num = topic_data['num']
                                content = topic_data['content']
                                topic_strings.append(f"Тема №{topic_num} - {content}")
                        except (ValueError, TypeError):
                            pass

                    topics_text = "; ".join(topic_strings)

                subj_cell = ws.cell(row=row, column=4, value=topics_text)
                subj_cell.border = thin_border
                subj_cell.alignment = left_alignment

                # 🔹 ОПРЕДЕЛЯЕМ ПРЕПОДАВАТЕЛЯ
                if subject_code.lower() == 'med':
                    teacher_name = str(plan.med_teacher) if plan.med_teacher else ''
                else:
                    teacher_name = str(plan.teacher) if plan.teacher else ''

                ws.cell(row=row, column=5, value=teacher_name).border = thin_border

                location = getattr(plan, 'location', '') or ''
                ws.cell(row=row, column=6, value=location).border = thin_border

                ws.row_dimensions[row].height = 30
                row += 1

        current_date += timedelta(days=1)

    # Настройка ширины колонок
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 80
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 20

    # === 🔹 ИМЯ ФАЙЛА С АДРЕСОМ ===
    t_name = _sanitize(str(plan.teacher)) if plan.teacher else "Без_преподавателя"

    location_raw = getattr(plan, 'location', '') or 'без_адреса'
    location_clean = _sanitize(location_raw)
    if not location_clean:
        location_clean = 'без_адреса'

    sched_val = getattr(plan, 'schedule_type', None) or getattr(plan.group, 'schedule_type', None)
    sched_map = {
        'even': 'четная', 'odd': 'нечетная', 'evening': 'вечерняя',
        '1': 'четная', '2': 'нечетная', '3': 'вечерняя',
        'четная': 'четная', 'нечетная': 'нечетная', 'вечерняя': 'вечерняя'
    }
    sched_name = sched_map.get(str(sched_val).lower().strip(), "четная")

    filename_base = f"Расписание_гр{_sanitize(group_num)}_{t_name}_{time_type}_{sched_name}_{location_clean}"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    encoded_filename = quote(f"{filename_base}.xlsx")
    response['Content-Disposition'] = f'attachment; filename*=UTF-8\'\'{encoded_filename}'

    ascii_filename = f"{filename_base}.xlsx".encode('ascii', 'ignore').decode('ascii')
    response['Content-Disposition'] += f'; filename="{ascii_filename}"'

    wb.save(response)
    return response